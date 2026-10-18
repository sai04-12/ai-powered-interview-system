"""Resume Document Parsing and Structured Analysis Service."""
import io
import os
import re
import uuid
import zipfile
from typing import Any, Dict, List, Optional, Set, Tuple
from pydantic import ValidationError
import pypdf
import docx
from app.core.config import settings
from app.models.resume import ExtractionStatus
from app.schemas.resume import (
    EducationItem,
    ExperienceItem,
    ProjectItem,
    ResumeStructuredData,
)

# Common Technical Skills & Technologies Taxonomy (normalized lowercase keys)
SKILLS_TAXONOMY: Dict[str, str] = {
    # Languages
    "python": "Python",
    "javascript": "JavaScript",
    "typescript": "TypeScript",
    "dart": "Dart",
    "java": "Java",
    "c++": "C++",
    "c#": "C#",
    "golang": "Go",
    "go": "Go",
    "rust": "Rust",
    "ruby": "Ruby",
    "php": "PHP",
    "swift": "Swift",
    "kotlin": "Kotlin",
    "scala": "Scala",
    "sql": "SQL",
    "r": "R",
    "html": "HTML",
    "css": "CSS",
    # Frameworks & Libraries
    "flutter": "Flutter",
    "fastapi": "FastAPI",
    "django": "Django",
    "flask": "Flask",
    "react": "React",
    "react native": "React Native",
    "next.js": "Next.js",
    "vue": "Vue.js",
    "angular": "Angular",
    "node.js": "Node.js",
    "express": "Express.js",
    "spring": "Spring",
    "spring boot": "Spring Boot",
    ".net": ".NET",
    "asp.net": "ASP.NET",
    "tensorflow": "TensorFlow",
    "pytorch": "PyTorch",
    "scikit-learn": "Scikit-Learn",
    "pandas": "Pandas",
    "numpy": "NumPy",
    "opencv": "OpenCV",
    # Databases & Storage
    "postgresql": "PostgreSQL",
    "postgres": "PostgreSQL",
    "mysql": "MySQL",
    "sqlite": "SQLite",
    "mongodb": "MongoDB",
    "redis": "Redis",
    "elasticsearch": "Elasticsearch",
    "dynamodb": "DynamoDB",
    "cassandra": "Cassandra",
    "firebase": "Firebase",
    "supabase": "Supabase",
    # DevOps, Cloud & Tools
    "docker": "Docker",
    "kubernetes": "Kubernetes",
    "aws": "AWS",
    "amazon web services": "AWS",
    "gcp": "GCP",
    "google cloud": "GCP",
    "azure": "Azure",
    "git": "Git",
    "github": "GitHub",
    "gitlab": "GitLab",
    "ci/cd": "CI/CD",
    "jenkins": "Jenkins",
    "terraform": "Terraform",
    "ansible": "Ansible",
    "linux": "Linux",
    "rest api": "REST API",
    "rest": "REST API",
    "graphql": "GraphQL",
    "grpc": "gRPC",
    "microservices": "Microservices",
    "kafka": "Kafka",
    "rabbitmq": "RabbitMQ",
    "celery": "Celery",
    # Practices & Methodologies
    "agile": "Agile",
    "scrum": "Scrum",
    "tdd": "TDD",
    "test driven development": "TDD",
    "unit testing": "Unit Testing",
    "oop": "OOP",
}

# Injection Attempt Patterns (Untrusted input detection)
INJECTION_PATTERNS = [
    re.compile(r"ignore\s+(?:all\s+)?previous\s+instructions", re.IGNORECASE),
    re.compile(r"system\s+prompt", re.IGNORECASE),
    re.compile(r"you\s+are\s+now\s+a", re.IGNORECASE),
    re.compile(r"grant\s+(?:all\s+)?access", re.IGNORECASE),
    re.compile(r"bypass\s+(?:all\s+)?verification", re.IGNORECASE),
    re.compile(r"give\s+(?:this\s+candidate|me)\s+(?:a\s+)?100%", re.IGNORECASE),
]


class ResumeParsingError(Exception):
    """Custom sanitized exception for resume upload and parsing failures."""
    pass


class ResumeService:
    """Service handling secure file validation, storage, text extraction, and analysis."""

    @staticmethod
    def validate_and_save_file(
        filename: str,
        content: bytes,
        upload_dir: str = settings.UPLOAD_DIR,
        max_size_mb: int = settings.MAX_UPLOAD_SIZE_MB,
    ) -> Tuple[str, str, int]:
        """
        Validate file size, extension, and magic byte signatures; store with unique server-generated key.
        Returns: (storage_key, mime_type, file_size)
        """
        # 1. Size Validation
        file_size = len(content)
        if file_size == 0:
            raise ResumeParsingError("Uploaded file is empty (0 bytes).")
        
        max_bytes = max_size_mb * 1024 * 1024
        if file_size > max_bytes:
            raise ResumeParsingError(f"File size exceeds maximum allowed limit of {max_size_mb} MB.")

        # 2. Extension Check
        clean_filename = os.path.basename(filename)
        ext = os.path.splitext(clean_filename)[1].lower()
        if ext not in [".pdf", ".docx"]:
            raise ResumeParsingError(f"Unsupported file format '{ext}'. Only PDF and DOCX files are permitted.")

        # 3. Magic Bytes / File Signature Validation
        mime_type = ""
        if ext == ".pdf":
            # PDF magic bytes: %PDF-
            if not content.startswith(b"%PDF-"):
                raise ResumeParsingError("Invalid file signature: uploaded file is not a valid PDF document.")
            mime_type = "application/pdf"
        elif ext == ".docx":
            # DOCX is a zip archive starting with PK\x03\x04
            if not (content.startswith(b"PK\x03\x04") or content.startswith(b"PK\x05\x06")):
                raise ResumeParsingError("Invalid file signature: uploaded file is not a valid DOCX document.")
            # Verify zip package integrity and Word document signature
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as zf:
                    namelist = zf.namelist()
                    if not any("word/document.xml" in name or "[Content_Types].xml" in name for name in namelist):
                        raise ResumeParsingError("Invalid DOCX format: missing Word document components.")
            except zipfile.BadZipFile:
                raise ResumeParsingError("Corrupted DOCX archive: file cannot be unpacked.")
            mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        # 4. Generate Safe Server-Side Storage Key & Prevent Path Traversal
        storage_key = f"{uuid.uuid4().hex}{ext}"
        os.makedirs(upload_dir, exist_ok=True)
        resolved_upload_dir = os.path.abspath(upload_dir)
        destination_path = os.path.abspath(os.path.join(resolved_upload_dir, storage_key))

        # Enforce that destination stays within upload_dir (Path Traversal defense)
        if not destination_path.startswith(resolved_upload_dir):
            raise ResumeParsingError("Invalid file path: path traversal detected.")

        with open(destination_path, "wb") as f:
            f.write(content)

        return storage_key, mime_type, file_size

    @staticmethod
    def extract_text_from_file(file_path: str, file_type: str) -> Tuple[str, bool]:
        """
        Extract clean text from PDF or DOCX file safely.
        Returns: (extracted_text, requires_ocr)
        """
        if not os.path.exists(file_path):
            raise ResumeParsingError("File not found on storage server.")

        text = ""
        requires_ocr = False

        try:
            if "pdf" in file_type.lower():
                reader = pypdf.PdfReader(file_path)
                if reader.is_encrypted:
                    try:
                        # Attempt to decrypt with empty password
                        reader.decrypt("")
                    except Exception:
                        raise ResumeParsingError("PDF file is password-protected or encrypted.")

                extracted_pages = []
                for page in reader.pages:
                    page_text = page.extract_text() or ""
                    extracted_pages.append(page_text.strip())

                text = "\n\n".join(extracted_pages).strip()

                # If text is too short or empty, document is likely a scanned image requiring OCR
                if len(text) < 30:
                    requires_ocr = True
                    text = ""

            elif "docx" in file_type.lower() or "word" in file_type.lower():
                doc = docx.Document(file_path)
                paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                # Also extract text from tables
                table_texts = []
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            table_texts.append(" | ".join(row_cells))
                
                all_chunks = paragraphs + table_texts
                text = "\n\n".join(all_chunks).strip()

            else:
                raise ResumeParsingError(f"Unsupported document type: {file_type}")

        except ResumeParsingError:
            raise
        except Exception as e:
            raise ResumeParsingError(f"Document extraction failed: {str(e)}")

        return text, requires_ocr

    @classmethod
    def sanitize_untrusted_text(cls, raw_text: str) -> str:
        """Neutralize malicious injection instructions embedded inside uploaded documents."""
        sanitized = raw_text
        for pattern in INJECTION_PATTERNS:
            sanitized = pattern.sub("[FILTERED_POTENTIAL_INSTRUCTION]", sanitized)
        return sanitized

    @classmethod
    def analyze_resume_text(cls, raw_text: str, requires_ocr: bool = False) -> ResumeStructuredData:
        """
        Deterministic, rule-based structured analysis of extracted resume text.
        Never invents qualifications, preserves uncertainty, handles missing sections gracefully.
        """
        if requires_ocr or not raw_text.strip():
            return ResumeStructuredData(
                professional_summary=None,
                skills=[],
                education=[],
                work_experience=[],
                projects=[],
                certifications=[],
                total_experience_years=None,
                technologies=[],
                missing_or_ambiguous=["No extractable text found; document requires optical character recognition (OCR)."],
                requires_ocr=True,
            )

        clean_text = cls.sanitize_untrusted_text(raw_text)
        missing_or_ambiguous: List[str] = []

        # 1. Skills & Technologies Extraction
        extracted_skills_set: Set[str] = set()
        for key, display_name in SKILLS_TAXONOMY.items():
            # Word-boundary or punctuation-boundary regex matching
            escaped_key = re.escape(key)
            pattern = rf"(?:\b|(?<=[^a-zA-Z0-9])){escaped_key}(?:\b|(?=[^a-zA-Z0-9]))"
            if re.search(pattern, clean_text, re.IGNORECASE):
                extracted_skills_set.add(display_name)

        extracted_skills = sorted(list(extracted_skills_set))
        if not extracted_skills:
            missing_or_ambiguous.append("No standardized technical skills identified in document text.")

        # 2. Professional Summary Extraction
        summary: Optional[str] = None
        summary_patterns = [
            r"(?:summary|professional\s+summary|profile|about\s+me|objective)[:\s\n]+([^\n]+(?:\n[^\n]+){1,5})",
        ]
        for pat in summary_patterns:
            match = re.search(pat, clean_text, re.IGNORECASE)
            if match:
                candidate_summary = match.group(1).strip()
                if len(candidate_summary) > 20:
                    summary = candidate_summary[:500]
                    break

        # 3. Education Extraction
        education: List[EducationItem] = []
        degree_patterns = [
            (r"(?:bachelor(?:'s)?|b\.?s\.?|b\.?tech|b\.?e\.?)\s+(?:in|of)?\s+([A-Za-z\s]+)", "Bachelor's"),
            (r"(?:master(?:'s)?|m\.?s\.?|m\.?tech|m\.?e\.?|mba)\s+(?:in|of)?\s+([A-Za-z\s]+)", "Master's"),
            (r"(?:ph\.?d|doctorate)\s+(?:in|of)?\s+([A-Za-z\s]+)", "PhD"),
            (r"(?:associate(?:'s)?|diploma)\s+(?:in|of)?\s+([A-Za-z\s]+)", "Associate's/Diploma"),
        ]
        for pat, degree_type in degree_patterns:
            matches = re.finditer(pat, clean_text, re.IGNORECASE)
            for m in matches:
                field = m.group(1).strip().split("\n")[0][:60]
                education.append(EducationItem(degree=degree_type, field_of_study=field))

        if not education:
            missing_or_ambiguous.append("Education background not explicitly recognized.")

        # 4. Work Experience & Duration Extraction
        work_experience: List[ExperienceItem] = []
        # Find sections starting with Experience / Work History
        exp_header = re.search(r"(?:work\s+experience|experience|employment\s+history)[:\s\n]", clean_text, re.IGNORECASE)
        if exp_header:
            exp_text = clean_text[exp_header.end():exp_header.end() + 1500]
            # Detect year ranges like 2018 - 2022, 2021 - Present
            year_matches = list(re.finditer(r"\b(20\d{2}|19\d{2})\s*[-–to]+\s*(20\d{2}|present|current)\b", exp_text, re.IGNORECASE))
            for ym in year_matches[:4]:
                start_year = int(ym.group(1))
                end_str = ym.group(2).lower()
                end_year = 2026 if "present" in end_str or "current" in end_str else int(end_str)
                duration = max(0.5, float(end_year - start_year))
                work_experience.append(
                    ExperienceItem(
                        title="Professional Role",
                        duration_years=duration,
                        description=f"Active tenure {ym.group(0)}"
                    )
                )

        total_experience: Optional[float] = None
        if work_experience:
            total_experience = round(sum(exp.duration_years or 0.0 for exp in work_experience), 1)
        else:
            # Check for direct mentions like "5+ years of experience"
            years_mention = re.search(r"(\d+(?:\.\d+)?)\+?\s*years?\s+(?:of\s+)?experience", clean_text, re.IGNORECASE)
            if years_mention:
                try:
                    total_experience = float(years_mention.group(1))
                except ValueError:
                    pass

        if total_experience is None:
            missing_or_ambiguous.append("Total years of experience could not be reliably verified.")

        # 5. Projects Extraction
        projects: List[ProjectItem] = []
        project_header = re.search(r"(?:projects|key\s+projects)[:\s\n]", clean_text, re.IGNORECASE)
        if project_header:
            proj_text = clean_text[project_header.end():project_header.end() + 1000]
            lines = [line.strip() for line in proj_text.split("\n") if line.strip() and len(line.strip()) > 5]
            for line in lines[:3]:
                # Identify project title / description
                projects.append(ProjectItem(name=line[:60], description=line[:150]))

        # 6. Certifications
        certifications: List[str] = []
        cert_matches = re.findall(
            r"(?:aws\s+certified[^\n,]+|certified\s+kubernetes[^\n,]+|pmp|comptia[^\n,]+|cisco\s+certified[^\n,]+)",
            clean_text,
            re.IGNORECASE,
        )
        for cert in cert_matches[:5]:
            clean_cert = cert.strip()
            if clean_cert and clean_cert not in certifications:
                certifications.append(clean_cert)

        return ResumeStructuredData(
            professional_summary=summary,
            skills=extracted_skills,
            education=education,
            work_experience=work_experience,
            projects=projects,
            certifications=certifications,
            total_experience_years=total_experience,
            technologies=extracted_skills,
            missing_or_ambiguous=missing_or_ambiguous,
            requires_ocr=False,
        )

    @classmethod
    def process_resume(
        cls,
        filename: str,
        content: bytes,
        upload_dir: str = settings.UPLOAD_DIR,
        max_size_mb: int = settings.MAX_UPLOAD_SIZE_MB,
    ) -> Tuple[str, str, int, str, Optional[str], Dict[str, Any]]:
        """
        End-to-end pipeline: validation, persistence, extraction, and structured analysis.
        Returns: (storage_key, mime_type, file_size, extraction_status, raw_text, structured_dict)
        """
        storage_key, mime_type, file_size = cls.validate_and_save_file(
            filename=filename,
            content=content,
            upload_dir=upload_dir,
            max_size_mb=max_size_mb,
        )
        file_path = os.path.join(upload_dir, storage_key)

        raw_text, requires_ocr = cls.extract_text_from_file(file_path, mime_type)
        analysis = cls.analyze_resume_text(raw_text, requires_ocr=requires_ocr)

        status = ExtractionStatus.FAILED.value if requires_ocr or not raw_text else ExtractionStatus.COMPLETED.value
        return (
            storage_key,
            mime_type,
            file_size,
            status,
            raw_text if raw_text else None,
            analysis.model_dump(),
        )
